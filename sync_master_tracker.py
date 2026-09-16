"""
Unified Master Vehicles Tracker
-------------------------------
Maintains the single master tracking spreadsheet in the parent workspace directory:
d:/Temp/yt automation/master_vehicles_tracker.xlsx

Combines real-time data across:
1. Video Generation (youtube-automation)
2. Video Editing (video-editor)
3. YouTube Uploading (youtube-uploader)

Columns:
- No. / ID
- Vehicle Name
- Category (Car / Motorcycle)
- Scale
- Color Scheme
- Generated (COMPLETED / PENDING / FAILED)
- Edited (EDITED / PENDING)
- Uploaded (UPLOADED / PENDING)
- Overall Stage
- Edited Video Location
- YouTube Short Link
"""

import os
import json
from datetime import datetime
from typing import Dict, Any, List

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def update_master_tracker(root_dir: str = None) -> str:
    """
    Reads JSON records from generation, editing, and uploading pipelines
    and synchronizes d:/Temp/yt automation/master_vehicles_tracker.xlsx.
    """
    if not root_dir:
        root_dir = os.path.dirname(os.path.abspath(__file__))

    xlsx_path = os.path.join(root_dir, "master_vehicles_tracker.xlsx")
    vehicles_json = os.path.join(root_dir, "youtube-automation", "data", "vehicles.json")
    edited_json = os.path.join(root_dir, "video-editor", "edited_vehicles.json")
    uploaded_json = os.path.join(root_dir, "youtube-uploader", "uploaded_videos.json")

    # 1. Load Vehicles Database (Generation)
    vehicles = []
    if os.path.exists(vehicles_json):
        try:
            with open(vehicles_json, "r", encoding="utf-8") as f:
                vehicles = json.load(f)
        except Exception:
            vehicles = []

    # 2. Load Edited Videos
    edited_map = {}
    if os.path.exists(edited_json):
        try:
            with open(edited_json, "r", encoding="utf-8") as f:
                for item in json.load(f):
                    if item.get("edit_status") == "EDITED":
                        edited_map[item.get("id")] = item
        except Exception:
            edited_map = {}

    # 3. Load Uploaded Videos
    uploaded_map = {}
    if os.path.exists(uploaded_json):
        try:
            with open(uploaded_json, "r", encoding="utf-8") as f:
                for item in json.load(f):
                    if item.get("upload_status") == "UPLOADED":
                        uploaded_map[item.get("id")] = item
        except Exception:
            uploaded_map = {}

    if not HAS_OPENPYXL:
        print("[-] openpyxl not installed, skipping Excel generation.")
        return xlsx_path

    wb = openpyxl.Workbook()

    # Style System
    font_family = "Segoe UI"
    title_font = Font(name=font_family, size=16, bold=True, color="1F497D")
    subtitle_font = Font(name=font_family, size=10, italic=True, color="595959")
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")

    kpi_fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    data_font = Font(name=font_family, size=10)
    link_font = Font(name=font_family, size=10, color="0000EE", underline="single")

    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    # Status fills and fonts
    fill_success = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid") # Soft Green
    font_success = Font(name=font_family, size=10, bold=True, color="274E13")

    fill_pending = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Soft Yellow
    font_pending = Font(name=font_family, size=10, bold=True, color="7F6000")

    fill_failed = PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid") # Soft Red/Orange
    font_failed = Font(name=font_family, size=10, bold=True, color="783F04")

    # -------------------------------------------------------------
    # Sheet 1: Master Overview
    # -------------------------------------------------------------
    ws = wb.active
    ws.title = "Master Automation Tracker"

    ws["A1"] = "YouTube Automation & Production Pipeline - Master Tracker"
    ws["A1"].font = title_font
    ws["A2"] = f"Real-time status across Generation, Editing, and Uploading | Last updated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws["A2"].font = subtitle_font

    total_vehicles = len(vehicles) if vehicles else 1
    gen_count = len([v for v in vehicles if v.get("status") == "COMPLETED"])
    edit_count = len(edited_map)
    upload_count = len(uploaded_map)
    pending_gen_count = len([v for v in vehicles if v.get("status") == "PENDING"])

    # KPI Summary Cards (Rows 4-5)
    cards = [
        ("B4", "B5", "TOTAL VEHICLES", str(len(vehicles)), "1F497D"),
        ("D4", "D5", "1. GENERATED", f"{gen_count} ({gen_count * 100 // total_vehicles}%)", "137333"),
        ("F4", "F5", "2. EDITED", f"{edit_count} ({edit_count * 100 // total_vehicles}%)", "274E13"),
        ("H4", "H5", "3. UPLOADED", f"{upload_count} ({upload_count * 100 // total_vehicles}%)", "B71C1C"),
        ("J4", "J5", "REMAINING GEN", f"{pending_gen_count}", "7F6000"),
    ]

    for top, bot, label, val, col_hex in cards:
        ws[top] = label
        ws[top].font = Font(name=font_family, size=9, bold=True, color="595959")
        ws[top].alignment = center_align
        ws[top].fill = kpi_fill
        ws[top].border = thin_border

        ws[bot] = val
        ws[bot].font = Font(name=font_family, size=13, bold=True, color=col_hex)
        ws[bot].alignment = center_align
        ws[bot].fill = kpi_fill
        ws[bot].border = thin_border

    headers = [
        "No. / ID", "Vehicle Name", "Category", "Scale", "Color Scheme",
        "Generated", "Edited", "Uploaded", "Overall Stage", "Edited Video Path", "YouTube Short Link"
    ]

    start_row = 7
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    for r_idx, v in enumerate(vehicles, start=start_row + 1):
        vid = v.get("id")
        gen_status = v.get("status", "PENDING")
        is_edited = vid in edited_map
        is_uploaded = vid in uploaded_map

        edit_status = "EDITED" if is_edited else ("PENDING" if gen_status == "COMPLETED" else "-")
        upload_status = "UPLOADED" if is_uploaded else ("PENDING" if is_edited else "-")

        # Overall stage determination
        if is_uploaded:
            overall_stage = "Published on YouTube 🚀"
        elif is_edited:
            overall_stage = "Edited (Ready to Upload) ✨"
        elif gen_status == "COMPLETED":
            overall_stage = "Generated (Pending Edit) 🎬"
        elif gen_status == "FAILED":
            overall_stage = "Generation Failed ❌"
        else:
            overall_stage = "Remaining to Generate ⏳"

        edited_path = edited_map[vid].get("edited_video_path", "-") if is_edited else "-"
        yt_link = uploaded_map[vid].get("youtube_url", "-") if is_uploaded else "-"

        row_data = [
            vid,
            v.get("vehicle_name"),
            v.get("type", "").capitalize(),
            v.get("scale"),
            v.get("color_scheme"),
            gen_status,
            edit_status,
            upload_status,
            overall_stage,
            edited_path,
            yt_link
        ]

        for c_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = center_align if c_idx in [1, 3, 4, 6, 7, 8, 9] else left_align

            # Color Highlights for Generated (col 6)
            if c_idx == 6:
                if gen_status == "COMPLETED":
                    cell.fill = fill_success; cell.font = font_success
                elif gen_status == "PENDING":
                    cell.fill = fill_pending; cell.font = font_pending
                else:
                    cell.fill = fill_failed; cell.font = font_failed

            # Color Highlights for Edited (col 7)
            elif c_idx == 7:
                if edit_status == "EDITED":
                    cell.fill = fill_success; cell.font = font_success
                elif edit_status == "PENDING":
                    cell.fill = fill_pending; cell.font = font_pending
                else:
                    cell.font = Font(name=font_family, size=10, color="888888")

            # Color Highlights for Uploaded (col 8)
            elif c_idx == 8:
                if upload_status == "UPLOADED":
                    cell.fill = fill_success; cell.font = font_success
                elif upload_status == "PENDING":
                    cell.fill = fill_pending; cell.font = font_pending
                else:
                    cell.font = Font(name=font_family, size=10, color="888888")

            # Hyperlink for YouTube Short
            elif c_idx == 11 and isinstance(val, str) and val.startswith("http"):
                cell.font = link_font
                cell.hyperlink = val

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    try:
        wb.save(xlsx_path)
        print(f"[+] Master tracker updated: {xlsx_path}")
    except PermissionError:
        fallback_path = os.path.join(root_dir, "master_vehicles_tracker_live.xlsx")
        try:
            wb.save(fallback_path)
            print(f"[!] Note: {os.path.basename(xlsx_path)} is open in Excel. Saved live copy to {os.path.basename(fallback_path)}!")
        except Exception:
            print(f"[!] Warning: Could not overwrite {xlsx_path} (file is open in Excel). Please close it.")
    except Exception as e:
        print(f"[-] Error saving master tracker: {e}")

    return xlsx_path


if __name__ == "__main__":
    update_master_tracker()
