import os
import json
from typing import List, Dict, Any, Optional
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.utils.logger import get_logger

logger = get_logger("excel_tracker")


class ExcelTracker:
    """
    Maintains and synchronizes a professional Excel tracker for Cars and Bikes automation.
    Tracks status (COMPLETED, PENDING, FAILED), video file paths, timestamps, and specs.
    """
    def __init__(self, excel_path: str = "data/vehicles_tracker.xlsx"):
        if not os.path.isabs(excel_path):
            base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
            excel_path = os.path.join(base_dir, excel_path)
        self.excel_path = excel_path

    def sync_from_vehicles_list(self, vehicles: List[Dict[str, Any]]) -> str:
        """
        Creates/updates the Excel tracking workbook from the vehicle records.
        """
        os.makedirs(os.path.dirname(self.excel_path), exist_ok=True)
        wb = openpyxl.Workbook()

        # Styles
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        
        status_completed_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
        status_completed_font = Font(name="Segoe UI", size=10, bold=True, color="274E13")

        status_pending_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        status_pending_font = Font(name="Segoe UI", size=10, bold=True, color="7F6000")

        status_failed_fill = PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid")
        status_failed_font = Font(name="Segoe UI", size=10, bold=True, color="783F04")

        data_font = Font(name="Segoe UI", size=10)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        # -------------------------------------------------------------
        # Sheet 1: Master Tracker
        # -------------------------------------------------------------
        ws_all = wb.active
        ws_all.title = "All Vehicles Tracker"

        headers = [
            "ID", "Vehicle Name", "Category", "Color Scheme", "Scale",
            "Status", "Completed Timestamp", "Output Video Path"
        ]
        ws_all.append(headers)

        for col_num, _ in enumerate(headers, 1):
            cell = ws_all.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        for row_idx, v in enumerate(vehicles, start=2):
            status = v.get("status", "PENDING")
            row_data = [
                v.get("id"),
                v.get("vehicle_name"),
                v.get("type", "").capitalize(),
                v.get("color_scheme", ""),
                v.get("scale", ""),
                status,
                v.get("completed_date") or "-",
                v.get("output_video_path") or "-"
            ]
            ws_all.append(row_data)

            # Apply cell formatting
            for col_idx in range(1, len(headers) + 1):
                c = ws_all.cell(row=row_idx, column=col_idx)
                c.font = data_font
                c.border = thin_border
                c.alignment = center_align if col_idx in [1, 3, 5, 6, 7] else left_align

                if col_idx == 6:  # Status column
                    if status == "COMPLETED":
                        c.fill = status_completed_fill
                        c.font = status_completed_font
                    elif status == "PENDING":
                        c.fill = status_pending_fill
                        c.font = status_pending_font
                    else:
                        c.fill = status_failed_fill
                        c.font = status_failed_font

        # Auto-adjust column widths
        for col in ws_all.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_all.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # -------------------------------------------------------------
        # Sheet 2: Completed Videos Only
        # -------------------------------------------------------------
        ws_completed = wb.create_sheet(title="Completed Videos")
        comp_headers = [
            "ID", "Vehicle Name", "Category", "Color Scheme",
            "Scale", "Completed Date", "Video File Path"
        ]
        ws_completed.append(comp_headers)

        for col_num, _ in enumerate(comp_headers, 1):
            cell = ws_completed.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = PatternFill(start_color="274E13", end_color="274E13", fill_type="solid")
            cell.alignment = center_align

        completed_vehicles = [v for v in vehicles if v.get("status") == "COMPLETED"]
        for row_idx, v in enumerate(completed_vehicles, start=2):
            row_data = [
                v.get("id"),
                v.get("vehicle_name"),
                v.get("type", "").capitalize(),
                v.get("color_scheme", ""),
                v.get("scale", ""),
                v.get("completed_date") or "-",
                v.get("output_video_path") or "-"
            ]
            ws_completed.append(row_data)

            for col_idx in range(1, len(comp_headers) + 1):
                c = ws_completed.cell(row=row_idx, column=col_idx)
                c.font = data_font
                c.border = thin_border
                c.alignment = center_align if col_idx in [1, 3, 5, 6] else left_align

        for col in ws_completed.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_completed.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(self.excel_path)
        logger.info(f"✓ Excel tracker updated successfully at: {self.excel_path} ({len(completed_vehicles)} completed / {len(vehicles)} total)")
        return self.excel_path

    def sync_from_json(self, json_path: str = "data/vehicles.json") -> str:
        """Reads JSON database and updates Excel file."""
        if not os.path.isabs(json_path):
            base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
            json_path = os.path.join(base_dir, json_path)

        if os.path.exists(json_path):
            with open(json_path, "r", encoding="utf-8") as f:
                vehicles = json.load(f)
            return self.sync_from_vehicles_list(vehicles)
        return ""
