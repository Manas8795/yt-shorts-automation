"""
Popular Cars USA & Canada Top 50 Dedicated Tracker
--------------------------------------------------
Maintains the dedicated tracking spreadsheet in the parent workspace directory:
d:/Temp/yt automation/Popular_Cars_USA_Canada_tracker.xlsx

Tracks all 50 popular USA & Canada cars across:
1. Video Generation (youtube-automation/popular_cars_usa_canada)
2. Video Editing (video-editor)
3. YouTube Uploading (youtube-uploader)
"""

import os
import json
from datetime import datetime
from typing import Dict, Any, List

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def update_usa_cars_tracker(root_dir: str = None) -> str:
    """
    Synchronizes d:/Temp/yt automation/Popular_Cars_USA_Canada_tracker.xlsx
    with live data from popular cars USA/Canada pipeline.
    """
    if not root_dir:
        root_dir = os.path.dirname(os.path.abspath(__file__))

    xlsx_path = os.path.join(root_dir, "Popular_Cars_USA_Canada_tracker.xlsx")
    cars_json = os.path.join(root_dir, "youtube-automation", "popular_cars_usa_canada", "vehicles.json")
    edited_json = os.path.join(root_dir, "video-editor", "edited_vehicles.json")
    uploaded_json = os.path.join(root_dir, "youtube-uploader", "uploaded_videos.json")

    # 1. Load Cars Database
    cars = []
    if os.path.exists(cars_json):
        try:
            with open(cars_json, "r", encoding="utf-8") as f:
                cars = json.load(f)
        except Exception:
            cars = []

    # 2. Load Edited Videos (filtered to Popular Cars USA & Canada)
    edited_map = {}
    if os.path.exists(edited_json):
        try:
            with open(edited_json, "r", encoding="utf-8") as f:
                for item in json.load(f):
                    v_src = item.get("source_excel", "")
                    v_path = item.get("edited_video_path", "")
                    if v_src == "Popular_Cars_USA_Canada_Top_50" or "Popular_Cars" in v_path:
                        v_name = item.get("vehicle_name", "").lower()
                        if v_name:
                            edited_map[v_name] = item
        except Exception:
            edited_map = {}

    # 3. Load Uploaded Videos (filtered to Popular Cars USA & Canada)
    uploaded_map = {}
    if os.path.exists(uploaded_json):
        try:
            with open(uploaded_json, "r", encoding="utf-8") as f:
                for item in json.load(f):
                    v_src = item.get("source_excel", "")
                    v_path = item.get("local_video_path", "")
                    if v_src == "Popular_Cars_USA_Canada_Top_50" or "Popular_Cars" in v_path:
                        v_name = item.get("vehicle_name", "").lower()
                        if v_name:
                            uploaded_map[v_name] = item
        except Exception:
            uploaded_map = {}

    if not HAS_OPENPYXL:
        print("[-] openpyxl not installed, skipping Excel creation.")
        return xlsx_path

    wb = openpyxl.Workbook()

    # Style definitions
    font_family = "Segoe UI"
    title_font = Font(name=font_family, size=16, bold=True, color="003366")
    subtitle_font = Font(name=font_family, size=10, italic=True, color="595959")
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")

    kpi_fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    data_font = Font(name=font_family, size=10)
    link_font = Font(name=font_family, size=10, color="0000EE", underline="single")

    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    thick_bottom = Border(bottom=Side(style="medium", color="003366"))

    ws = wb.active
    ws.title = "USA Canada Top 50"
    ws.views.sheetView[0].showGridLines = True

    # Title Block
    ws["B2"] = "POPULAR CARS USA & CANADA TOP 50 — SHORTS PIPELINE TRACKER"
    ws["B2"].font = title_font
    ws["B3"] = f"Generated & Maintained by Automated Pipeline | Last Updated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws["B3"].font = subtitle_font

    # Calculate Metrics
    total_cars = len(cars)
    gen_done = sum(1 for c in cars if c.get("status") == "COMPLETED")
    gen_failed = sum(1 for c in cars if c.get("status") == "FAILED")
    gen_pending = sum(1 for c in cars if c.get("status") in ["PENDING", None])

    edit_done = 0
    up_done = 0
    for c in cars:
        c_name = c.get("vehicle_name", "").lower()
        if c_name in edited_map and edited_map[c_name].get("edit_status") == "EDITED":
            edit_done += 1
        if c_name in uploaded_map and uploaded_map[c_name].get("upload_status") == "UPLOADED":
            up_done += 1

    gen_pct = (gen_done / total_cars * 100) if total_cars else 0
    edit_pct = (edit_done / total_cars * 100) if total_cars else 0
    up_pct = (up_done / total_cars * 100) if total_cars else 0

    # Summary KPI Cards
    summary_headers = [
        ("B5", "TOTAL CARS"),
        ("D5", "1. GENERATED"),
        ("F5", "2. EDITED"),
        ("H5", "3. UPLOADED")
    ]
    summary_values = [
        ("B6", f"{total_cars}"),
        ("D6", f"{gen_done} ({gen_pct:.0f}%)"),
        ("F6", f"{edit_done} ({edit_pct:.0f}%)"),
        ("H6", f"{up_done} ({up_pct:.0f}%)")
    ]

    for cell_ref, text in summary_headers:
        ws[cell_ref] = text
        ws[cell_ref].font = Font(name=font_family, size=9, bold=True, color="595959")
        ws[cell_ref].alignment = center_align

    for cell_ref, val in summary_values:
        ws[cell_ref] = val
        ws[cell_ref].font = Font(name=font_family, size=14, bold=True, color="003366")
        ws[cell_ref].alignment = center_align

    # Table Headers (Row 8)
    headers = [
        ("B", "No. / ID", 10),
        ("C", "Car Name", 28),
        ("D", "Generated", 16),
        ("E", "Edited", 16),
        ("F", "Uploaded", 16),
        ("G", "Overall Stage", 24),
        ("H", "Color Scheme & Appearance", 65),
        ("I", "Scale", 10),
        ("J", "Local Video Location", 50),
        ("K", "YouTube Short Link", 38)
    ]

    for col_letter, title, width in headers:
        cell = ws[f"{col_letter}8"]
        cell.value = title
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[8].height = 28

    # Fill Vehicle Data
    start_row = 9
    for idx, c in enumerate(cars, start=0):
        current_row = start_row + idx
        ws.row_dimensions[current_row].height = 22
        is_even = (idx % 2 == 0)
        row_bg = "FFFFFF" if is_even else "F9FAFB"
        row_fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")

        vid_id = c.get("id", idx + 1)
        v_name = c.get("vehicle_name", "")
        norm_name = v_name.lower()
        color_scheme = c.get("color_scheme", "")
        scale = c.get("scale", "1:18")

        # Stage 1: Generation
        gen_status = c.get("status", "PENDING").upper()
        if gen_status == "COMPLETED":
            gen_display = "COMPLETED"
            gen_color = "0E6251"
            gen_bg = "D4EFDF"
        elif gen_status == "FAILED":
            gen_display = "FAILED"
            gen_color = "922B21"
            gen_bg = "FADBD8"
        else:
            gen_display = "PENDING"
            gen_color = "7D6608"
            gen_bg = "FCF3CF"

        # Stage 2: Editing
        edit_rec = edited_map.get(norm_name)
        if edit_rec and edit_rec.get("edit_status") == "EDITED":
            edit_display = "EDITED"
            edit_color = "145A32"
            edit_bg = "D5F5E3"
            video_location = edit_rec.get("edited_video_path", "")
        else:
            edit_display = "PENDING"
            edit_color = "566573"
            edit_bg = "EAECEE"
            video_location = c.get("output_video_path", "")

        # Stage 3: Uploading
        upload_rec = uploaded_map.get(norm_name)
        if upload_rec and upload_rec.get("upload_status") == "UPLOADED":
            upload_display = "UPLOADED"
            upload_color = "1B4F72"
            upload_bg = "D4E6F1"
            yt_link = upload_rec.get("video_url") or "https://studio.youtube.com/channel/videos/short"
        else:
            upload_display = "PENDING"
            upload_color = "566573"
            upload_bg = "EAECEE"
            yt_link = ""

        # Overall Status
        if upload_display == "UPLOADED":
            stage_display = "Published on YouTube"
            stage_color = "1B4F72"
            stage_bg = "E8F8F5"
        elif edit_display == "EDITED":
            stage_display = "Ready for Upload"
            stage_color = "196F3D"
            stage_bg = "EAFAF1"
        elif gen_display == "COMPLETED":
            stage_display = "Ready for Editing"
            stage_color = "B7950B"
            stage_bg = "FEF9E7"
        else:
            stage_display = "Awaiting Generation"
            stage_color = "5D6D7E"
            stage_bg = "EBEDEF"

        # Cells population
        cells_data = [
            ("B", vid_id, center_align, data_font, row_fill),
            ("C", v_name, left_align, Font(name=font_family, size=10, bold=True), row_fill),
            ("D", gen_display, center_align, Font(name=font_family, size=9, bold=True, color=gen_color), PatternFill(start_color=gen_bg, end_color=gen_bg, fill_type="solid")),
            ("E", edit_display, center_align, Font(name=font_family, size=9, bold=True, color=edit_color), PatternFill(start_color=edit_bg, end_color=edit_bg, fill_type="solid")),
            ("F", upload_display, center_align, Font(name=font_family, size=9, bold=True, color=upload_color), PatternFill(start_color=upload_bg, end_color=upload_bg, fill_type="solid")),
            ("G", stage_display, center_align, Font(name=font_family, size=9, bold=True, color=stage_color), PatternFill(start_color=stage_bg, end_color=stage_bg, fill_type="solid")),
            ("H", color_scheme, left_align, data_font, row_fill),
            ("I", scale, center_align, data_font, row_fill),
            ("J", video_location, left_align, Font(name=font_family, size=8, color="333333"), row_fill),
            ("K", yt_link, left_align, link_font if yt_link else data_font, row_fill)
        ]

        for col_letter, val, align, fnt, fll in cells_data:
            c_cell = ws[f"{col_letter}{current_row}"]
            c_cell.value = val
            c_cell.alignment = align
            c_cell.font = fnt
            c_cell.fill = fll
            c_cell.border = thin_border

    wb.save(xlsx_path)
    return xlsx_path


if __name__ == "__main__":
    out = update_usa_cars_tracker()
    print(f"[+] Synced Popular Cars USA & Canada tracker: {out}")
