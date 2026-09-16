"""
Upload Tracker for YouTube Shorts
---------------------------------
Maintains tracking records for all uploaded Shorts in JSON and Excel.
"""

import os
import json
from datetime import datetime
from typing import List, Dict, Any, Optional

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class UploadTracker:
    def __init__(self, script_dir: Optional[str] = None):
        if not script_dir:
            script_dir = os.path.dirname(os.path.abspath(__file__))
        self.script_dir = script_dir
        self.json_path = os.path.join(script_dir, "uploaded_videos.json")
        self.xlsx_path = os.path.join(script_dir, "upload_tracker.xlsx")
        self.records: Dict[int, Dict[str, Any]] = {}
        self.load()

    def load(self):
        if os.path.exists(self.json_path):
            try:
                with open(self.json_path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    self.records_list = data
                    self.records = {}
                    for item in data:
                        vid_id = item.get("id")
                        v_name = item.get("vehicle_name", "").strip().lower()
                        v_path = os.path.normpath(item.get("local_video_path", ""))
                        v_src = item.get("source_excel") or ("100_popular_cars_in_India" if "100_popular_cars" in v_path else "master_vehicles_tracker")
                        item["source_excel"] = v_src

                        # Key by composite key and name
                        self.records[f"{v_src}_{vid_id}"] = item
                        if v_name:
                            self.records[f"name_{v_name}"] = item
                        if v_path:
                            self.records[f"path_{v_path}"] = item
                        # Keep fallback id key for single tracker backwards compatibility
                        if vid_id not in self.records:
                            self.records[vid_id] = item
            except Exception:
                self.records = {}
                self.records_list = []
        else:
            self.records = {}
            self.records_list = []

    def is_uploaded(
        self,
        vid_id: int,
        vehicle_name: Optional[str] = None,
        source_excel: Optional[str] = None,
        video_path: Optional[str] = None
    ) -> bool:
        if source_excel:
            comp_key = f"{source_excel}_{vid_id}"
            if comp_key in self.records:
                return self.records[comp_key].get("upload_status") == "UPLOADED"
            if video_path:
                norm = os.path.normpath(video_path)
                for item in getattr(self, 'records_list', []):
                    if item.get("source_excel") == source_excel and os.path.basename(item.get("local_video_path", "")) == os.path.basename(norm):
                        return item.get("upload_status") == "UPLOADED"
            return False

        if vehicle_name and f"name_{vehicle_name.strip().lower()}" in self.records:
            return self.records[f"name_{vehicle_name.strip().lower()}"].get("upload_status") == "UPLOADED"
        if video_path:
            norm = os.path.normpath(video_path)
            for k, v in self.records.items():
                if isinstance(k, str) and k.startswith("path_") and os.path.basename(norm) in k:
                    return v.get("upload_status") == "UPLOADED"
        # Fallback to ID check only if source_excel wasn't provided
        if vid_id in self.records:
            return self.records[vid_id].get("upload_status") == "UPLOADED"
        return False

    def record_upload(
        self,
        vid_id: int,
        vehicle_name: str,
        category: str,
        title: str,
        youtube_url: str,
        visibility: str,
        local_video_path: str,
        source_excel: Optional[str] = None
    ):
        v_path = os.path.normpath(local_video_path) if local_video_path else ""
        v_src = source_excel or ("100_popular_cars_in_India" if "100_popular_cars" in v_path else "master_vehicles_tracker")
        rec = {
            "id": vid_id,
            "vehicle_name": vehicle_name,
            "category": category,
            "source_excel": v_src,
            "title": title,
            "youtube_url": youtube_url,
            "visibility": visibility.upper(),
            "upload_status": "UPLOADED",
            "uploaded_date": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "local_video_path": local_video_path
        }
        # Update in-memory structures
        self.records[f"{v_src}_{vid_id}"] = rec
        if vehicle_name:
            self.records[f"name_{vehicle_name.strip().lower()}"] = rec
        if v_path:
            self.records[f"path_{v_path}"] = rec
        self.records[vid_id] = rec

        # Deduplicate and append to records_list
        updated_list = []
        found = False
        for item in self.records_list:
            if item.get("vehicle_name", "").strip().lower() == vehicle_name.strip().lower() or (item.get("id") == vid_id and item.get("source_excel") == v_src):
                updated_list.append(rec)
                found = True
            else:
                updated_list.append(item)
        if not found:
            updated_list.append(rec)
        self.records_list = updated_list
        self.save()

    def save(self):
        # 1. Save JSON
        sorted_records = sorted(self.records_list, key=lambda x: (x.get("source_excel", ""), x.get("id", 0)))
        with open(self.json_path, "w", encoding="utf-8") as f:
            json.dump(sorted_records, f, indent=2)

        # 2. Save Excel
        if not HAS_OPENPYXL:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Uploaded Shorts"

        font_family = "Segoe UI"
        header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="B71C1C", end_color="B71C1C", fill_type="solid") # YouTube Red
        data_font = Font(name=font_family, size=10)
        link_font = Font(name=font_family, size=10, color="0000EE", underline="single")
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        headers = [
            "No. / ID", "Vehicle Name", "Category", "Short Title",
            "YouTube URL", "Visibility", "Upload Status", "Uploaded Date", "Local Video File"
        ]
        ws.append(headers)

        for col_num, _ in enumerate(headers, 1):
            c = ws.cell(row=1, column=col_num)
            c.font = header_font
            c.fill = header_fill
            c.alignment = center_align

        for r_idx, r in enumerate(sorted_records, start=2):
            yt_url = r.get("youtube_url", "-")
            row_data = [
                r.get("id"),
                r.get("vehicle_name"),
                r.get("category", "").capitalize(),
                r.get("title"),
                yt_url,
                r.get("visibility"),
                r.get("upload_status"),
                r.get("uploaded_date"),
                r.get("local_video_path")
            ]
            ws.append(row_data)

            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=r_idx, column=col_idx)
                c.border = thin_border
                c.alignment = center_align if col_idx in [1, 3, 6, 7, 8] else left_align

                if col_idx == 5 and yt_url.startswith("http"):
                    c.font = link_font
                    c.hyperlink = yt_url
                else:
                    c.font = data_font

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(self.xlsx_path)
